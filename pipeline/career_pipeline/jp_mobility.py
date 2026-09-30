"""Japan: job changes and job ranks → web/public/data/jp/mobility.json

雇用動向調査 令和7年（厚生労働省, e-Stat Excel, no key）
  第13-2表  転職入職者（前職も雇用者、入職前1年以内に就業）の賃金変動区分の構成比 × 性 × 年齢
  参考表2    転職入職率 × 性 × 年齢
賃金構造基本統計調査 令和7年 役職（雇用期間の定めのない労働者、企業規模10人以上）
  役職第2表  部長級・課長級・係長級・職長級・非役職 × 性 × 学歴 × 年齢（労働者数・所定内給与額・年間賞与）
  役職第1表  役職計（上の4区分に含まれない役職を含む）

Output:
  payChange   {"age|sex": [n_thousand, up10, upLess10, same, downLess10, down10, unknown]}  (% of job changers)
  changeRate  {"age|sex": 転職入職率 %}
  ranks       {"age|sex|edu": {"bucho"|"kacho"|"kakari"|"shokucho"|"other"|"none": [workers, annualPay|null]}}
"""

from __future__ import annotations

import re
import time
import unicodedata
from collections import defaultdict

import openpyxl
import requests
import xlrd

from .common import RAW_DIR, WEB_DATA_DIR, write_json
from .jp_wage_census import EDU, SEX, age_code, merge, num, parse_layout_a

RAW = RAW_DIR / "jp" / "mobility"
DOWNLOAD = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={}&fileKind=4"
FILES = {
    "000040489954": "koyo_000040489954.xls",  # 第13-2表 賃金変動
    "000040490249": "koyo_000040490249.xls",  # 参考表2 労働移動率
    "000040420884": "chin_000040420884.xlsx",  # 役職第1表 産業計
    "000040421113": "chin_000040421113.xlsx",  # 役職第2表
}
KOYO_SEX = {"計": "*", "男": "M", "女": "F"}  # 雇用動向調査 labels the all-sexes row 「計」
RANKS = {"部長級": "bucho", "課長級": "kacho", "係長級": "kakari", "職長級": "shokucho", "非役職": "none"}
EDU_ORDER = ["学歴計", "中学", "高校", "専門学校", "高専・短大", "大学", "大学院", "不明"]


def fetch() -> None:
    RAW.mkdir(parents=True, exist_ok=True)
    for sid, name in FILES.items():
        path = RAW / name
        if not path.exists():
            print(f"downloading {sid}")
            r = requests.get(DOWNLOAD.format(sid), headers={"User-Agent": "Mozilla/5.0 (career-map data pipeline)"}, timeout=120)
            r.raise_for_status()
            path.write_bytes(r.content)
            time.sleep(1)


def norm(s) -> str:
    return re.sub(r"\s", "", unicodedata.normalize("NFKC", str(s)))


def koyo_age(label: str) -> str | None:
    t = norm(label)
    if t == "計":
        return "*"
    if t.endswith("以上"):
        return "65-69" if t.startswith("65") else None
    if t.startswith("19歳以下"):
        return "18-19"
    return age_code(t)


def pay_change() -> dict[str, list]:
    s = xlrd.open_workbook(RAW / FILES["000040489954"]).sheet_by_index(0)
    out: dict[str, list] = {}
    sex = None
    started = False
    for r in range(s.nrows):
        block = str(s.cell_value(r, 3))
        if "（前職）" in block:
            if started:
                break  # only the first block: 経験1年未満, (前職)計 → (現職)計
            started = True
            continue
        if not started:
            continue
        if norm(s.cell_value(r, 1)) in ("ＧＴ", "GT"):
            sex = KOYO_SEX.get(norm(s.cell_value(r, 2)))
        age = "*" if norm(s.cell_value(r, 1)) in ("ＧＴ", "GT") else koyo_age(s.cell_value(r, 2))
        if sex is None or age is None:
            continue
        vals = [s.cell_value(r, c) for c in (3, 18, 19, 20, 22, 23, 24)]
        if any(not isinstance(v, float) for v in vals):
            continue
        out[f"{age}|{sex}"] = [round(v, 1) for v in vals]
    return out


def change_rate() -> dict[str, float]:
    s = xlrd.open_workbook(RAW / FILES["000040490249"]).sheet_by_index(0)
    out: dict[str, float] = {}
    sex = None
    for r in range(9, s.nrows):
        if str(s.cell_value(r, 1)).strip() and norm(s.cell_value(r, 1)) not in ("ＧＴ", "GT"):
            break  # the next block (一般労働者) starts
        if norm(s.cell_value(r, 1)) in ("ＧＴ", "GT"):
            sex = KOYO_SEX.get(norm(s.cell_value(r, 2)))
            continue
        age = koyo_age(s.cell_value(r, 2))
        rate = s.cell_value(r, 10)
        if sex is not None and age and isinstance(rate, float):
            out[f"{age}|{sex}"] = rate
    return out


def rank_rows(ws):
    """役職第2表 sheet → {(sex, edu, age): (所定内, 賞与, 労働者数)} for 勤続年数計."""
    out = {}
    sex, edu = "*", None
    for row in ws.iter_rows(min_row=13, values_only=True):
        label = row[2]
        if label is None:
            continue
        age = "*"
        for part in str(label).split("\n"):
            tok = norm(part)
            if tok in SEX:
                sex = SEX[tok]
            elif (a := age_code(tok)) is not None:
                age = a
            elif tok:
                edu = tok
        s, b, n = (num(x) for x in row[3:6])
        if age != "70+" and n is not None:
            out[(sex, edu, age)] = (s, b, n)
    return out


def ranks() -> dict[str, dict[str, list]]:
    wb = openpyxl.load_workbook(RAW / FILES["000040421113"], read_only=True, data_only=True)
    by_rank = {code: rank_rows(wb[f"(10人以上){name}"]) for name, code in RANKS.items()}
    total_rows = parse_layout_a(
        openpyxl.load_workbook(RAW / FILES["000040420884"], read_only=True, data_only=True)["産業計(役職計)"],
        label_col=3, first_row=13,
    )
    ranked_total = {(r.sex, r.group, r.age): r.workers for r in total_rows}

    # Merge Wage Census 学歴 into the site's education levels (worker-weighted)
    groups: dict[tuple[str, str, str], dict[str, list]] = defaultdict(lambda: defaultdict(list))
    for code, rows in by_rank.items():
        for (sex, edu_raw, age), (s, b, n) in rows.items():
            edu = EDU.get(edu_raw or "")
            if edu is not None:
                groups[(sex, edu, age)][code].append((s, b, n))
    for (sex, edu_raw, age), n in ranked_total.items():
        edu = EDU.get(edu_raw or "")
        if edu is not None:
            groups[(sex, edu, age)]["ranked_total"].append((None, None, n))

    out: dict[str, dict[str, list]] = {}
    for (sex, edu, age), parts in groups.items():
        cell: dict[str, list] = {}
        for code, items in parts.items():
            n = sum(x[2] for x in items)
            paid = [x for x in items if x[0] is not None and x[1] is not None and x[2] > 0]
            w = sum(x[2] for x in paid)
            annual = round((12 * sum(x[0] * x[2] for x in paid) / w + sum(x[1] * x[2] for x in paid) / w) * 1000, -3) if w else None
            cell[code] = [round(n * 10), None if annual is None else int(annual)]
        if "ranked_total" not in cell or "none" not in cell:
            continue
        named = sum(cell[c][0] for c in ("bucho", "kacho", "kakari", "shokucho") if c in cell)
        cell["other"] = [max(cell.pop("ranked_total")[0] - named, 0), None]
        if sum(v[0] for v in cell.values()) >= 1000:
            out[f"{age}|{sex}|{edu}"] = cell
    return out


def build() -> None:
    data = {
        "source": {
            "payChange": {"ja": "厚生労働省「令和7年雇用動向調査」第13-2表", "en": "MHLW, Survey on Employment Trends 2025, table 13-2",
                          "url": "https://www.e-stat.go.jp/stat-search/files?toukei=00450073&tstat=000001012468"},
            "ranks": {"ja": "厚生労働省「令和7年賃金構造基本統計調査」役職第1表・第2表", "en": "MHLW, Basic Survey on Wage Structure 2025, job-rank tables 1-2",
                      "url": "https://www.e-stat.go.jp/stat-search/files?toukei=00450091&tstat=000001011429"},
        },
        "payChange": pay_change(),
        "changeRate": change_rate(),
        "ranks": ranks(),
    }
    write_json(WEB_DATA_DIR / "jp" / "mobility.json", data)
    print(f"  payChange {len(data['payChange'])}, changeRate {len(data['changeRate'])}, ranks {len(data['ranks'])}")


if __name__ == "__main__":
    fetch()
    build()
