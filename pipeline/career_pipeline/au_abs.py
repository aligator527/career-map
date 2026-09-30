"""Australia: ABS Employee Earnings and Hours (EEH) May 2025 → wage-distribution cells.

Population: full-time non-managerial employees paid at the adult rate (the ABS headline population
for earnings distributions), weekly total cash earnings × 52.

  EEH 63060DO008  deciles by occupation (ANZSCO major) × sex, by age, by state      → method 0
  EEH 63060DO006  means by age × occupation major, age × sex, age × state          → method 3
  EEH 63060DO013  means by occupation (ANZSCO 3-digit minor group) × sex            → method 3

Method 3 cells take their mean from the table and the shape (p10…p90 relative to the mean) from the
closest cell with published deciles (occupation major × sex, age, or state).

EEH does not cross states with occupation or sex, and has no education. For those cells the ABS
Characteristics of Employment survey, August 2025 (6337.0 tables 4a and 5; median and mean weekly
earnings of full-time employees by state × occupation major × sex, and state × qualification × sex)
supplies *relative* differences that are applied to the closest EEH cell, e.g.

  region s, occupation k, sex g:  EEH national (k, g) × CoE (s, k, g) / CoE (Australia, k, g)
  education e, sex g:             EEH national (g)    × CoE (Australia, e, g) / CoE (Australia, all, g)

keeping the EEH level and shape (method 3). CoE estimates with a relative standard error above 25%
are dropped.
"""

from __future__ import annotations

import json
import time
from pathlib import Path

import openpyxl
import requests

from .common import RAW_DIR, WEB_DATA_DIR, write_json, write_national, write_region_summary

EEH_URL = "https://www.abs.gov.au/statistics/labour/earnings-and-working-conditions/employee-earnings-and-hours-australia/may-2025"
COE_URL = "https://www.abs.gov.au/statistics/labour/earnings-and-working-conditions/characteristics-employment-australia/aug-2025"
EEH_CUBES = ["006", "008", "013"]
COE_TABLES = ["04a", "05"]
COE_MONTH = ("Aug-25", "Aug-2025")
AU_RAW = RAW_DIR / "au"
OUT_DIR = WEB_DATA_DIR / "au"
LABELS_JA = json.loads((Path(__file__).parent / "labels" / "au_occupations_ja.json").read_text())

WEEKS = 52
METHOD_DIRECT = 0
METHOD_SHAPE = 3
MAX_RSE = 25.0  # CoE estimates less reliable than this are not used

# ASGS 2021 state and territory codes
REGIONS = {
    "1": ("New South Wales", "ニューサウスウェールズ州", "NSW"),
    "2": ("Victoria", "ビクトリア州", "VIC"),
    "3": ("Queensland", "クイーンズランド州", "QLD"),
    "4": ("South Australia", "南オーストラリア州", "SA"),
    "5": ("Western Australia", "西オーストラリア州", "WA"),
    "6": ("Tasmania", "タスマニア州", "TAS"),
    "7": ("Northern Territory", "北部準州", "NT"),
    "8": ("Australian Capital Territory", "オーストラリア首都特別地域", "ACT"),
}
STATE_CODE = {en.lower(): code for code, (en, _, _) in REGIONS.items()}

# ANZSCO 2022 major groups
MAJORS = {
    "1": "Managers",
    "2": "Professionals",
    "3": "Technicians and trades workers",
    "4": "Community and personal service workers",
    "5": "Clerical and administrative workers",
    "6": "Sales workers",
    "7": "Machinery operators and drivers",
    "8": "Labourers",
}
MAJOR_CODE = {en.lower(): code for code, en in MAJORS.items()}

AGE = {
    "20 years and under": "18-20", "18 to 20 years": "18-20", "21 to 24 years": "21-24",
    "25 to 34 years": "25-34", "35 to 44 years": "35-44", "45 to 54 years": "45-54",
    "55 to 64 years": "55-64", "65 years and over": "65-69",
}
AGES = ["18-20", "21-24", "25-34", "35-44", "45-54", "55-64", "65-69"]

# CoE "Level of highest non-school qualification" → education codes used by the web app
EDU = {
    "without non-school qualification": "secondary",
    "advanced diploma or diploma": "short_tertiary",
    "bachelor degree": "bachelor",
    "postgraduate degree": "graduate",
}
EDUCATIONS = ["secondary", "short_tertiary", "bachelor", "graduate"]

SEX_BLOCK = {"males": "M", "females": "F", "persons": "*"}
MEAN_WEEKLY = "AVERAGE WEEKLY TOTAL CASH EARNINGS ($)"
PCT_ROWS = ("10th", "25th", "50th", "75th", "90th")

UA = {"User-Agent": "Mozilla/5.0 (career-map data pipeline)"}


def eeh_path(cube: str) -> Path:
    return AU_RAW / f"63060DO{cube}_202505.xlsx"


def coe_path(table: str) -> Path:
    return AU_RAW / f"63370_Table{table}.xlsx"


def fetch() -> None:
    AU_RAW.mkdir(parents=True, exist_ok=True)
    jobs = [(eeh_path(c), f"{EEH_URL}/{eeh_path(c).name}") for c in EEH_CUBES]
    jobs += [(coe_path(t), f"{COE_URL}/{coe_path(t).name}") for t in COE_TABLES]
    for path, url in jobs:
        if path.exists():
            continue
        print(f"downloading {url}")
        r = requests.get(url, headers=UA, timeout=600)
        r.raise_for_status()
        path.write_bytes(r.content)
        time.sleep(1)


# ---------------------------------------------------------------- EEH data cubes


def num(v) -> float | None:
    return float(v) if isinstance(v, (int, float)) else None


def eeh_table(cube: str, sheet: str) -> dict[tuple[str, ...], dict[str, float | None]]:
    """ESTIMATE block of an EEH table as {(measure, sub-block, row label): {column label: value}}.

    Section headers are rows with a label and no values: upper-case ones are measures ("NUMBER OF
    EMPLOYEES ('000)"), others sub-blocks ("Males"). Reading stops at the STANDARD ERROR block.
    """
    ws = openpyxl.load_workbook(eeh_path(cube), read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    cols = [str(c).strip() if c is not None else None for c in rows[4]]
    out: dict[tuple[str, ...], dict[str, float | None]] = {}
    measure, block = "", ""
    for row in rows[5:]:
        label = str(row[0]).strip() if row[0] is not None else ""
        if not label:
            continue
        if label.startswith("STANDARD ERROR") or label.startswith("©"):
            break
        values = {cols[i]: num(v) for i, v in enumerate(row) if i > 0 and i < len(cols) and cols[i]}
        if all(v is None for v in values.values()):
            if label.split("(")[0].isupper():
                measure, block = label, ""
            else:
                block = label
            continue
        out[(measure, block, label)] = values
    return out


def eeh_standard_errors(cube: str, sheet: str) -> dict[str, list]:
    """Rows of the standard-error table (same layout as the estimate table) by row label."""
    ws = openpyxl.load_workbook(eeh_path(cube), read_only=True, data_only=True)[sheet]
    return {str(r[0]).strip(): list(r) for r in ws.iter_rows(min_row=8, values_only=True) if r[0]}


def deciles(table, block: str, column: str) -> tuple[list[float], float] | None:
    """[p10, p25, p50, p75, p90] and mean (weekly) of one column of a deciles table."""
    rows = {label: vals for (_, b, label), vals in table.items() if b == block}
    q = []
    for prefix in PCT_ROWS:
        (label,) = [x for x in rows if x.startswith(prefix)]
        q.append(rows[label].get(column))
    (mean_label,) = [x for x in rows if x.startswith("Average")]
    mean = rows[mean_label].get(column)
    if mean is None or any(x is None for x in q):
        return None
    return q, mean


def totals(table, block: str) -> dict[str, float]:
    """'Total' row of a distribution table (number of employees, thousands) by column."""
    for (_, b, label), vals in table.items():
        if b == block and label == "Total":
            return {k: v for k, v in vals.items() if v is not None}
    return {}


# ---------------------------------------------------------------- Characteristics of Employment


def coe_extract(table: str) -> dict[tuple[str, str, str], dict[str, float]]:
    """Full-time median, mean and employees for every (state, category, sex) of a CoE table, August
    2025, all employees (with and without leave entitlements). Cached as JSON next to the workbook
    because the workbooks are 25–35 MB."""
    cache = AU_RAW / f"63370_Table{table}.fulltime.json"
    src = coe_path(table)
    if cache.exists() and cache.stat().st_mtime >= src.stat().st_mtime:
        raw = json.loads(cache.read_text())
        return {tuple(k.split("|")): v for k, v in raw.items()}
    ws = openpyxl.load_workbook(src, read_only=True, data_only=True).worksheets[2]
    out: dict[tuple[str, str, str], dict[str, float]] = {}
    for r in ws.iter_rows(min_row=9, values_only=True):
        if str(r[0]).strip() not in COE_MONTH or str(r[2]).strip() != "Total employees":
            continue
        cls, cat = str(r[4]).strip(), str(r[5]).strip()
        if cat != "Full-time" or cls not in ("Full-time or part-time status", "Mean earnings"):
            continue
        state = str(r[1]).strip().lower()
        category = str(r[3]).strip().lower()
        for sex, off in (("*", 6), ("M", 12), ("F", 18)):
            d = out.setdefault((state, category, sex), {})
            v, rse = num(r[off]), num(r[off + 1])
            if cls == "Mean earnings":
                d["mean"], d["mean_rse"] = v or 0.0, rse or 0.0
            else:
                d["median"], d["median_rse"] = v or 0.0, rse or 0.0
                d["n"] = (num(r[off + 4]) or 0.0) * 1000
    cache.write_text(json.dumps({"|".join(k): v for k, v in out.items()}))
    return out


def coe_ok(d: dict | None) -> bool:
    return bool(d) and d.get("median", 0) > 0 and 0 < d.get("median_rse", 0) <= MAX_RSE


# ---------------------------------------------------------------- cells


def annual(x: float) -> int:
    return max(100, int(round(x * WEEKS, -2)))


def direct_cell(n: float, q: list[float], mean: float) -> list:
    return [round(n), annual(mean)] + [annual(x) for x in q] + [METHOD_DIRECT]


def shaped_cell(n: float, mean: float, shape: list) -> list:
    """Cell with this weekly mean and the quantile/mean ratios of `shape` (an annual cell)."""
    ratio = [x / shape[1] for x in shape[2:7]]
    return [round(n), annual(mean)] + [annual(mean * r) for r in ratio] + [METHOD_SHAPE]


def scaled_cell(base: list, n: float, r_median: float, r_mean: float) -> list:
    """`base` (annual cell) with quantiles scaled by r_median and the mean by r_mean."""
    return ([round(n), max(100, int(round(base[1] * r_mean, -2)))]
            + [max(100, int(round(x * r_median, -2))) for x in base[2:7]] + [METHOD_SHAPE])


def ratio(num_: dict, den: dict) -> tuple[float, float]:
    r_med = num_["median"] / den["median"]
    ok_mean = num_.get("mean", 0) > 0 and den.get("mean", 0) > 0 and num_.get("mean_rse", 99) <= MAX_RSE
    return r_med, (num_["mean"] / den["mean"] if ok_mean else r_med)


def build() -> None:
    national: dict[str, list] = {}
    regions: dict[str, dict[str, list]] = {code: {} for code in REGIONS}

    # --- DO008: published deciles
    occ_q = eeh_table("008", "Table_8")
    occ_n = eeh_table("008", "Table_7")
    for block, sex in SEX_BLOCK.items():
        blk = block.capitalize()
        n = totals(occ_n, blk)
        for col in [c for c in next(iter(occ_q.values())).keys()]:
            d = deciles(occ_q, blk, col)
            if d is None:
                continue
            occ = "*" if col == "All occupations" else f"M{MAJOR_CODE[col.lower()]}"
            national[f"{occ}|*|{sex}|*"] = direct_cell(n.get(col, 0) * 1000, *d)

    age_q = eeh_table("008", "Table_2")
    age_n = totals(eeh_table("008", "Table_1"), "")
    for col, age in AGE.items():
        d = deciles(age_q, "", col) if col in next(iter(age_q.values())) else None
        if d:
            national[f"*|{age}|*|*"] = direct_cell(age_n.get(col, 0) * 1000, *d)

    state_q = eeh_table("008", "Table_14")
    state_n = totals(eeh_table("008", "Table_13"), "")
    for name, code in STATE_CODE.items():
        col = REGIONS[code][0]
        d = deciles(state_q, "", col)
        if d:
            regions[code]["*|*|*|*"] = direct_cell(state_n.get(col, 0) * 1000, *d)

    # --- DO006: means by age × sex, age × occupation, age × state (shape of the published deciles)
    t1 = eeh_table("006", "Table_1")
    for block, sex in SEX_BLOCK.items():
        if sex == "*":
            continue
        n = t1.get(("NUMBER OF EMPLOYEES ('000)", block.capitalize(), "All employees"), {})
        mean = t1.get((MEAN_WEEKLY, block.capitalize(), "All employees"), {})
        for col, age in AGE.items():
            if mean.get(col) and f"*|{age}|*|*" in national:
                national[f"*|{age}|{sex}|*"] = shaped_cell((n.get(col) or 0) * 1000, mean[col], national[f"*|{age}|*|*"])

    t2 = eeh_table("006", "Table_2")
    for (measure, block, label), vals in t2.items():
        if measure != MEAN_WEEKLY or block or label.lower() not in MAJOR_CODE:
            continue
        occ = f"M{MAJOR_CODE[label.lower()]}"
        n = t2.get(("NUMBER OF EMPLOYEES ('000)", "", label), {})
        for col, age in AGE.items():
            if vals.get(col) and (n.get(col) or 0) >= 1:
                national[f"{occ}|{age}|*|*"] = shaped_cell(n[col] * 1000, vals[col], national[f"{occ}|*|*|*"])

    t4 = eeh_table("006", "Table_4")
    for (measure, block, label), vals in t4.items():
        code = STATE_CODE.get(label.lower())
        if measure != MEAN_WEEKLY or block or code is None:
            continue
        n = t4.get(("NUMBER OF EMPLOYEES ('000)", "", label), {})
        for col, age in AGE.items():
            if vals.get(col) and (n.get(col) or 0) >= 1 and f"*|{age}|*|*" in national:
                regions[code][f"*|{age}|*|*"] = shaped_cell(n[col] * 1000, vals[col], national[f"*|{age}|*|*"])

    # --- DO013: means by ANZSCO 3-digit minor group × sex (shape of the major group × sex)
    occupations: dict[str, str] = {}
    ws = openpyxl.load_workbook(eeh_path("013"), read_only=True, data_only=True)["Table_1"]
    se = eeh_standard_errors("013", "Table_2")
    for r in ws.iter_rows(min_row=8, values_only=True):
        label = str(r[0] or "").strip()
        if not (label[:3].isdigit() and label[3:4] == " "):
            continue
        code, name = label[:3], label[4:].strip()
        for sex, col in (("M", 2), ("F", 8), ("*", 14)):  # average weekly total cash earnings
            mean, err = num(r[col]), num(se.get(label, [None] * 20)[col])
            if not mean or (err is not None and err / mean > MAX_RSE / 100):
                continue
            shape = national.get(f"M{code[0]}|*|{sex}|*")
            if shape:
                occupations[code] = name
                national[f"{code}|*|{sex}|*"] = shaped_cell(0, mean, shape)

    # --- CoE: relative differences by state × occupation × sex and by qualification
    occ_coe = coe_extract("04a")
    edu_coe = coe_extract("05")
    au = "australia"
    for code, (en, _, _) in REGIONS.items():
        state = en.lower()
        base_region = regions[code]["*|*|*|*"]
        tot_state = edu_coe.get((state, "total", "*"))
        for sex in ("*", "M", "F"):
            d = edu_coe.get((state, "total", sex))
            if sex != "*" and coe_ok(d) and coe_ok(tot_state):
                regions[code][f"*|*|{sex}|*"] = scaled_cell(base_region, d["n"], *ratio(d, tot_state))
            for label, edu in EDU.items():
                d = edu_coe.get((state, label, sex))
                if coe_ok(d) and coe_ok(tot_state):
                    regions[code][f"*|*|{sex}|{edu}"] = scaled_cell(base_region, d["n"], *ratio(d, tot_state))
            for major in MAJORS:
                d, ref = occ_coe.get((state, MAJORS[major].lower(), sex)), occ_coe.get((au, MAJORS[major].lower(), sex))
                base = national.get(f"M{major}|*|{sex}|*")
                if base and coe_ok(d) and coe_ok(ref):
                    regions[code][f"M{major}|*|{sex}|*"] = scaled_cell(base, d["n"], *ratio(d, ref))

    for sex in ("*", "M", "F"):
        ref = edu_coe.get((au, "total", sex))
        for label, edu in EDU.items():
            d = edu_coe.get((au, label, sex))
            if coe_ok(d) and coe_ok(ref):
                national[f"*|*|{sex}|{edu}"] = scaled_cell(national[f"*|*|{sex}|*"], d["n"], *ratio(d, ref))

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for f in OUT_DIR.glob("region-*.json"):
        f.unlink()
    write_national(OUT_DIR, national, lambda occ: occ[0])
    for code, cells in regions.items():
        write_json(OUT_DIR / f"region-{code}.json", cells)
    write_json(OUT_DIR / "meta.json", {
        "country": "AU",
        "currency": "AUD",
        "year": 2025,
        "wageDefinition": {
            "en": "Weekly total cash earnings (including overtime and bonuses paid in the pay period) of full-time "
                  "non-managerial employees paid at the adult rate, May 2025, × 52. Excludes managerial employees, "
                  "juniors, apprentices and the self-employed. State × occupation, state × sex and education "
                  "differences are from the ABS Characteristics of Employment survey (August 2025, all full-time employees).",
            "ja": "フルタイムの非管理職被用者（成人賃金適用者）の週当たり現金給与総額（残業代・その期間に支払われた賞与を含む。"
                  "2025年5月）を52倍して年換算。管理職被用者・年少者・見習い・自営業者を除く。州×職業・州×性別・学歴別の差は"
                  "ABS「雇用特性調査」（2025年8月、全フルタイム被用者）による。",
        },
        "nKind": "population",
        "source": {
            "id": "eeh2025",
            "name": {"en": "ABS, Employee Earnings and Hours, Australia, May 2025",
                     "ja": "オーストラリア統計局「被用者の所得と労働時間（EEH）」2025年5月"},
            "url": EEH_URL,
        },
        "regions": [{"code": k, "abbr": a, "label": {"en": en, "ja": ja}} for k, (en, ja, a) in REGIONS.items()],
        "occupationMajor": [{"code": f"M{k}", "label": {"en": v, "ja": LABELS_JA[f"M{k}"]}} for k, v in MAJORS.items()],
        "occupations": [{"code": k, "major": k[0], "label": {"en": v[:1].upper() + v[1:], "ja": LABELS_JA[k]}}
                        for k, v in sorted(occupations.items())],
        "ages": AGES,
        "educations": EDUCATIONS,
    })
    write_region_summary(OUT_DIR)
    print(f"  AU: {len(national)} national cells, {len(regions)} regions, {len(occupations)} occupations")


if __name__ == "__main__":
    fetch()
    build()
