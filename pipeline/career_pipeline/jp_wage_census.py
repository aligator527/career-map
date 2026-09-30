"""Japan: Basic Survey on Wage Structure (賃金構造基本統計調査) 令和7年 → wage-distribution cells.

The survey publishes only tabulations (no microdata), as Excel files on e-Stat that need no API key.
We combine three kinds of tables:

  * means per cell (monthly scheduled/contractual pay, annual bonus, worker count)
      職種×性×年齢, 職種×年齢, 職種×性, 職種大分類×性×年齢, 学歴×性×年齢, 都道府県×性×年齢, 都道府県×職種大分類×性
  * histograms of monthly scheduled pay per 性×学歴×年齢 (学歴 第3表)
  * published quantiles of monthly scheduled pay per 職種 (職種 第17〜19表)

and produce, per cell, the annual mean and p10..p90 of annual pay (see DESIGN.md §6.2):

  method 1 "histogram": quantiles from the tabulated histogram, scaled from monthly scheduled pay
                        to annual pay by the ratio of the cell means.
  method 2 "lognormal": lognormal with the cell's annual mean and a dispersion taken from the
                        occupation's published quantiles, adjusted for age.
"""

from __future__ import annotations

import json
import math
import re
import time
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import requests

from .common import AGE_BANDS, QUANTILES, RAW_DIR, WEB_DATA_DIR, write_json, write_region_summary

YEAR = 2025  # 令和7年
JP_RAW = RAW_DIR / "jp"
OUT_DIR = WEB_DATA_DIR / "jp"
LABELS = json.loads((Path(__file__).parent / "labels" / "jp_occupations.json").read_text())

DOWNLOAD_URL = "https://www.e-stat.go.jp/stat-search/file-download?statInfId={}&fileKind=4"
TABLE_PAGE = "https://www.e-stat.go.jp/stat-search/files?toukei=00450091&tstat=000001011429"

# statInfId of each e-Stat file (all 一般労働者, 民営, 企業規模10人以上)
T_OCC_SEX_AGE = "000040421122"  # 職種（特掲）×性×年齢 第7表
T_OCC_AGE = "000040421120"  # 職種（小分類）×年齢 第5表（男女計）
T_OCC_SEX = "000040421116"  # 職種（小分類）×性 第1表
T_MAJOR_SEX_AGE = "000040421121"  # 職種（大分類）×性×年齢 第6表
T_OCC_DIST = ["000040421132", "000040421133"]  # 職種（小分類）分布特性値 第17表（男女計）
T_OCC_SEX_DIST = ["000040421135", "000040421136"]  # 職種（特掲）×性 分布特性値 第19表
T_MAJOR_DIST = "000040421134"  # 職種（大分類）×性 分布特性値 第18表
T_EDU_AGE = "000040420843"  # 学歴×性×年齢 第1表
T_EDU_AGE_DIST = "000040420874"  # 学歴×性×年齢 所定内給与額階級 第3表
T_PREF_AGE = [f"0000404211{n}" for n in range(67, 91)]  # 都道府県別第1表（2県ずつ）
T_PREF_MAJOR = ["000040421191", "000040421192", "000040421193", "000040421194"]  # 都道府県別第2表
ALL_TABLES = [
    T_OCC_SEX_AGE, T_OCC_AGE, T_OCC_SEX, T_MAJOR_SEX_AGE, *T_OCC_DIST, *T_OCC_SEX_DIST,
    T_MAJOR_DIST, T_EDU_AGE, T_EDU_AGE_DIST, *T_PREF_AGE, *T_PREF_MAJOR,
]

# Cells whose estimated worker count (十人) is below this are not published
MIN_WORKERS_10 = 100

PREFS = [
    ("北海道", "Hokkaido"), ("青森", "Aomori"), ("岩手", "Iwate"), ("宮城", "Miyagi"), ("秋田", "Akita"),
    ("山形", "Yamagata"), ("福島", "Fukushima"), ("茨城", "Ibaraki"), ("栃木", "Tochigi"), ("群馬", "Gunma"),
    ("埼玉", "Saitama"), ("千葉", "Chiba"), ("東京", "Tokyo"), ("神奈川", "Kanagawa"), ("新潟", "Niigata"),
    ("富山", "Toyama"), ("石川", "Ishikawa"), ("福井", "Fukui"), ("山梨", "Yamanashi"), ("長野", "Nagano"),
    ("岐阜", "Gifu"), ("静岡", "Shizuoka"), ("愛知", "Aichi"), ("三重", "Mie"), ("滋賀", "Shiga"),
    ("京都", "Kyoto"), ("大阪", "Osaka"), ("兵庫", "Hyogo"), ("奈良", "Nara"), ("和歌山", "Wakayama"),
    ("鳥取", "Tottori"), ("島根", "Shimane"), ("岡山", "Okayama"), ("広島", "Hiroshima"), ("山口", "Yamaguchi"),
    ("徳島", "Tokushima"), ("香川", "Kagawa"), ("愛媛", "Ehime"), ("高知", "Kochi"), ("福岡", "Fukuoka"),
    ("佐賀", "Saga"), ("長崎", "Nagasaki"), ("熊本", "Kumamoto"), ("大分", "Oita"), ("宮崎", "Miyazaki"),
    ("鹿児島", "Kagoshima"), ("沖縄", "Okinawa"),
]
PREF_CODE = {ja: f"{i + 1:02d}" for i, (ja, _) in enumerate(PREFS)}
PREF_FULL = {"北海道": "北海道", "東京": "東京都", "京都": "京都府", "大阪": "大阪府"}

SEX = {"男女計": "*", "男": "M", "女": "F"}
# Wage Census 学歴 → common education level (shared with US)
EDU = {
    "学歴計": "*", "中学": "secondary", "高校": "secondary", "専門学校": "short_tertiary",
    "高専・短大": "short_tertiary", "大学": "bachelor", "大学院": "graduate",
}
EDU_IGNORED = {"不明"}

# Monthly scheduled-pay histogram bin edges, 千円 (学歴 第3表). The top bin is open; we close it at 1600.
BIN_EDGES = [50.0] + [100.0 + 20 * i for i in range(16)] + [450.0, 500.0, 550.0, 600.0, 700.0, 800.0, 900.0, 1000.0, 1200.0, 1600.0]
Z = {0.10: -1.2815516, 0.25: -0.6744898, 0.50: 0.0, 0.75: 0.6744898, 0.90: 1.2815516}

METHOD_HISTOGRAM = 1
METHOD_LOGNORMAL = 2

OCC_CODE = {o["ja"]: f"J{i + 1:03d}" for i, o in enumerate(LABELS["occupations"])}

# 日本標準産業分類 大分類 (sheets of 学歴 第1表) and 企業規模 column blocks of the Layout-A tables
INDUSTRY_EN = {
    "C": "Mining and quarrying", "D": "Construction", "E": "Manufacturing", "F": "Electricity, gas, heat and water",
    "G": "Information and communications", "H": "Transport and postal", "I": "Wholesale and retail",
    "J": "Finance and insurance", "K": "Real estate and leasing", "L": "Scientific research, professional and technical services",
    "M": "Accommodation and food services", "N": "Living-related services and amusement", "O": "Education and learning support",
    "P": "Medical, health care and welfare", "Q": "Compound services", "R": "Other services",
}
SIZES = {"1000+": (12, "1,000人以上", "1,000+ employees"), "100-999": (20, "100〜999人", "100–999 employees"),
         "10-99": (28, "10〜99人", "10–99 employees")}
MAJOR_CODE = {m["ja"]: f"M{k}" for k, m in LABELS["majors"].items()}


# ---------------------------------------------------------------- fetch

def fetch() -> None:
    JP_RAW.mkdir(parents=True, exist_ok=True)
    headers = {"User-Agent": "Mozilla/5.0 (career-map data pipeline)"}
    for sid in ALL_TABLES:
        if list(JP_RAW.rglob(f"*{sid}.xlsx")):
            continue
        print(f"downloading {sid}")
        r = requests.get(DOWNLOAD_URL.format(sid), headers=headers, timeout=120)
        r.raise_for_status()
        (JP_RAW / f"{sid}.xlsx").write_bytes(r.content)
        time.sleep(1)


def open_book(sid: str) -> openpyxl.Workbook:
    (path,) = JP_RAW.rglob(f"*{sid}.xlsx")
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


# ---------------------------------------------------------------- parsing helpers

def norm(s) -> str:
    return re.sub(r"\s", "", unicodedata.normalize("NFKC", str(s)))


def num(v) -> float | None:
    if v is None or isinstance(v, str):
        return None
    return float(v)


def age_code(tok: str) -> str | None:
    """'~19歳' → '18-19', '20~24歳' / '20~24' → '20-24', '70歳~' → '70+'. None if not an age label."""
    m = re.fullmatch(r"(~)?(\d+)(?:~(\d+))?歳?(~)?", tok)
    if not m:
        return None
    if m.group(4):
        return "70+"
    if m.group(1):
        return "18-19"
    return f"{m.group(2)}-{m.group(3)}" if m.group(3) else None


@dataclass
class MeanRow:
    """One row of a 'Layout A' wage table (企業規模計 block)."""
    sex: str
    group: str | None  # occupation / education name, or None
    age: str  # age band or "*"
    kimatte: float  # きまって支給する現金給与額 (千円/月)
    shoteinai: float  # 所定内給与額 (千円/月)
    bonus: float  # 年間賞与その他特別給与額 (千円/年)
    workers: float  # 労働者数 (十人)

    @property
    def annual(self) -> float:
        """Annual pay in yen: 12 × monthly contractual cash earnings + annual bonus."""
        return (12 * self.kimatte + self.bonus) * 1000


def parse_layout_a(ws, label_col: int, first_row: int = 11, sex: str = "*", value_col: int = 4) -> list[MeanRow]:
    rows: list[MeanRow] = []
    group: str | None = None
    for row in ws.iter_rows(min_row=first_row, values_only=True):
        label = row[label_col - 1]
        if label is None:
            continue
        age = "*"
        for part in str(label).split("\n"):
            tok = norm(part)
            if not tok:
                continue
            if tok in SEX:
                sex = SEX[tok]
            elif (a := age_code(tok)) is not None:
                age = a
            else:
                group = tok
        v = [num(x) for x in row[value_col - 1 : value_col + 7]]
        kimatte, shoteinai, bonus, workers = v[4], v[5], v[6], v[7]
        if None in (kimatte, shoteinai, bonus, workers) or age == "70+":
            continue
        rows.append(MeanRow(sex, group, age, kimatte, shoteinai, bonus, workers))
    return rows


QUANTILE_LABELS = {"第1・十分位数": 0.10, "第1・四分位数": 0.25, "中位数": 0.50, "第3・四分位数": 0.75, "第9・十分位数": 0.90}


def quantile_label(tok: str) -> float | None:
    for k, q in QUANTILE_LABELS.items():
        if tok.startswith(k):
            return q
    return None


def parse_transposed_quantiles(ws) -> dict[str, dict[float, float]]:
    """職種 分布特性値 tables: occupations as columns (the '区分' row), quantile rows. First block = 企業規模計."""
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r[1] is not None and norm(r[1]) == "区分")
    names = {c: norm(v) for c, v in enumerate(rows[head]) if v is not None and c >= 3}
    out: dict[str, dict[float, float]] = defaultdict(dict)
    for row in rows[head + 1 :]:
        q = next((q for v in row[:4] if v is not None and (q := quantile_label(norm(v))) is not None), None)
        if q is None:
            continue
        for c, name in names.items():
            if q not in out[name] and (x := num(row[c])) is not None:
                out[name][q] = x
    return out


def sigma_from_quantiles(qs: dict[float, float]) -> float | None:
    if qs.get(0.10) and qs.get(0.90):
        return math.log(qs[0.90] / qs[0.10]) / (Z[0.90] - Z[0.10])
    return None


# ---------------------------------------------------------------- histograms (学歴 第3表)

def parse_edu_histograms(ws) -> dict[tuple[str, str, str], list[float]]:
    """{(sex, edu, age): [26 bin counts in 十人]} for 企業規模計, edu in raw Wage Census names."""
    rows = list(ws.iter_rows(values_only=True))
    header = rows[8]
    age_cols: dict[int, str] = {}
    for c, v in enumerate(header):
        if v is None or c < 5:
            continue
        t = norm(v)
        age_cols[c] = "*" if t == "年齢計" else age_code(t)
    hist: dict[tuple[str, str, str], list[float]] = {}
    sex, edu = "*", None
    for row in rows[9:]:
        block, bin_label = row[3], row[2]
        if block is not None and bin_label is None and quantile_label(norm(block)) is None:
            for part in str(block).split("\n"):
                tok = norm(part)
                if tok in SEX:
                    sex = SEX[tok]
                elif tok:
                    edu = tok
            continue
        if bin_label is not None and edu is not None:
            for c, age in age_cols.items():
                if age and age != "70+":
                    hist.setdefault((sex, edu, age), []).append(num(row[c]) or 0.0)
    return hist


def hist_quantiles(counts: list[float]) -> dict[float, float] | None:
    total = sum(counts)
    if total <= 0:
        return None
    out: dict[float, float] = {}
    cum = 0.0
    it = iter(QUANTILES)
    q = next(it)
    for i, c in enumerate(counts):
        while q is not None and cum + c >= q * total and c > 0:
            lo, hi = BIN_EDGES[i], BIN_EDGES[i + 1]
            out[q] = lo + (hi - lo) * (q * total - cum) / c
            q = next(it, None)
        cum += c
    return out


# ---------------------------------------------------------------- cells

Cells = dict[str, list]


def key(occ="*", age="*", sex="*", edu="*") -> str:
    return f"{occ}|{age}|{sex}|{edu}"


def cell_histogram(workers: float, annual_mean: float, shoteinai_mean: float, mq: dict[float, float]) -> list:
    """Scale monthly scheduled-pay quantiles (千円) to annual pay using the ratio of means."""
    ratio = annual_mean / (12 * shoteinai_mean * 1000)
    qs = [mq[q] * 12 * 1000 * ratio for q in QUANTILES]
    return [round(workers * 10), int(round(annual_mean, -3))] + [int(round(x, -3)) for x in qs] + [METHOD_HISTOGRAM]


def cell_lognormal(workers: float, annual_mean: float, sigma: float) -> list:
    mu = math.log(annual_mean) - sigma**2 / 2
    qs = [math.exp(mu + sigma * Z[q]) for q in QUANTILES]
    return [round(workers * 10), int(round(annual_mean, -3))] + [int(round(x, -3)) for x in qs] + [METHOD_LOGNORMAL]


def merge(rows: list[MeanRow]) -> MeanRow:
    """Worker-weighted merge of several MeanRows (e.g. 中学 + 高校 → secondary)."""
    w = sum(r.workers for r in rows)
    avg = lambda f: sum(getattr(r, f) * r.workers for r in rows) / w  # noqa: E731
    r0 = rows[0]
    return MeanRow(r0.sex, r0.group, r0.age, avg("kimatte"), avg("shoteinai"), avg("bonus"), w)


def build() -> None:
    # --- distribution shape: histograms per sex × edu × age
    hist_raw = parse_edu_histograms(open_book(T_EDU_AGE_DIST)["産業計(規模計)"])
    hist: dict[tuple[str, str, str], list[float]] = defaultdict(lambda: [0.0] * 26)
    for (sex, edu, age), counts in hist_raw.items():
        if edu in EDU_IGNORED or norm(edu) not in EDU:
            continue
        k = (sex, EDU[norm(edu)], age)
        hist[k] = [a + b for a, b in zip(hist[k], counts)]
    hist_q = {k: q for k, c in hist.items() if (q := hist_quantiles(c))}
    sigma_sex_age = {
        (sex, age): sigma_from_quantiles(hist_q[(sex, "*", age)])
        for (sex, edu, age) in hist_q
        if edu == "*"
    }

    def age_adjust(sex: str, age: str) -> float:
        return sigma_sex_age[(sex, age)] / sigma_sex_age[(sex, "*")] if age != "*" else 1.0

    # --- occupation dispersion (all ages)
    occ_sigma: dict[tuple[str, str], float] = {}  # (occ code, sex) → sigma
    for sid in T_OCC_DIST:
        for name, qs in parse_transposed_quantiles(open_book(sid)["男女計"]).items():
            if name in OCC_CODE and (s := sigma_from_quantiles(qs)):
                occ_sigma[(OCC_CODE[name], "*")] = s
    for sid in T_OCC_SEX_DIST:
        wb = open_book(sid)
        for sheet, sex in (("男", "M"), ("女", "F")):
            for name, qs in parse_transposed_quantiles(wb[sheet]).items():
                if name in OCC_CODE and (s := sigma_from_quantiles(qs)):
                    occ_sigma[(OCC_CODE[name], sex)] = s
    wb = open_book(T_MAJOR_DIST)
    for sheet, sex in (("男女計", "*"), ("男", "M"), ("女", "F")):
        for name, qs in parse_transposed_quantiles(wb[sheet]).items():
            if name in MAJOR_CODE and (s := sigma_from_quantiles(qs)):
                occ_sigma[(MAJOR_CODE[name], sex)] = s
    major_of = {OCC_CODE[o["ja"]]: f"M{o['major']}" for o in LABELS["occupations"]}

    def occ_cell(code: str, r: MeanRow) -> list | None:
        sigma = occ_sigma.get((code, r.sex)) or occ_sigma.get((code, "*"))
        if sigma is None and code in major_of:
            sigma = occ_sigma.get((major_of[code], r.sex)) or occ_sigma.get((major_of[code], "*"))
        if sigma is None or r.workers < MIN_WORKERS_10:
            return None
        return cell_lognormal(r.workers, r.annual, sigma * age_adjust(r.sex, r.age))

    national: Cells = {}

    # --- occupations (detail): 第5表 occ×age (男女計), 第1表 occ×sex, 第7表 occ×sex×age (特掲)
    occ_rows = parse_layout_a(open_book(T_OCC_AGE)[open_book(T_OCC_AGE).sheetnames[0]], label_col=2)
    wb = open_book(T_OCC_SEX)
    occ_rows += parse_layout_a(wb[wb.sheetnames[0]], label_col=2)
    wb = open_book(T_OCC_SEX_AGE)
    for sheet, sex in (("男", "M"), ("女", "F")):
        occ_rows += parse_layout_a(wb[sheet], label_col=2, sex=sex)
    unknown = set()
    for r in occ_rows:
        code = OCC_CODE.get(r.group or "")
        if code is None:
            unknown.add(r.group)
            continue
        if (c := occ_cell(code, r)) is not None:
            national[key(code, r.age, r.sex)] = c
    print(f"  occupation names not in labels: {sorted(unknown)}")

    # --- occupations (major groups): 第6表
    wb = open_book(T_MAJOR_SEX_AGE)
    for sheet, sex in (("男女計", "*"), ("男", "M"), ("女", "F")):
        for r in parse_layout_a(wb[sheet], label_col=2, sex=sex):
            code = MAJOR_CODE.get(r.group or "")
            if code and (c := occ_cell(code, r)) is not None:
                national[key(code, r.age, r.sex)] = c

    # --- education × sex × age (and all-worker totals via 学歴計): 学歴 第1表
    edu_rows = parse_layout_a(open_book(T_EDU_AGE)["産業計"], label_col=3, first_row=13)
    grouped: dict[tuple[str, str, str], list[MeanRow]] = defaultdict(list)
    for r in edu_rows:
        edu = EDU.get(r.group or "")
        if edu is not None:
            grouped[(r.sex, edu, r.age)].append(r)
    national_means: dict[tuple[str, str], MeanRow] = {}
    for (sex, edu, age), rs in grouped.items():
        m = merge(rs)
        if edu == "*":
            national_means[(sex, age)] = m
        mq = hist_q.get((sex, edu, age))
        if mq and m.workers >= MIN_WORKERS_10:
            national[key(age=age, sex=sex, edu=edu)] = cell_histogram(m.workers, m.annual, m.shoteinai, mq)

    # --- prefectures: 都道府県別第1表（産業計 sheets）, 第2表（職種大分類）
    regions: dict[str, Cells] = defaultdict(dict)
    for sid in T_PREF_AGE:
        wb = open_book(sid)
        for sheet in wb.sheetnames:
            m = re.fullmatch(r"\((.+)\)産業計", norm(sheet))
            if not m:
                continue
            code = PREF_CODE[m.group(1)]
            for r in parse_layout_a(wb[sheet], label_col=3, first_row=12):
                nat = national_means.get((r.sex, r.age))
                mq = hist_q.get((r.sex, "*", r.age))
                if nat and mq and r.workers >= MIN_WORKERS_10:
                    # Regional shape = national shape for the same sex × age, rescaled to the regional mean
                    # (quantiles = national annual quantiles × regional mean / national mean)
                    c = cell_histogram(nat.workers, nat.annual, nat.shoteinai, mq)
                    c[0], c[1] = round(r.workers * 10), int(round(r.annual, -3))
                    c[2:7] = [int(round(x * r.annual / nat.annual, -3)) for x in c[2:7]]
                    regions[code][key(age=r.age, sex=r.sex)] = c
    for sid in T_PREF_MAJOR:
        ws = open_book(sid).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        pref_cols = {}
        for c, v in enumerate(rows[5]):
            t = norm(v) if v else ""
            name = re.sub(r"^\d+", "", t)
            if name in PREF_CODE:
                pref_cols[c + 1] = PREF_CODE[name]
        for col, code in pref_cols.items():
            for r in parse_layout_a(ws, label_col=2, first_row=10, value_col=col):
                mcode = MAJOR_CODE.get(r.group or "")
                if mcode and (c := occ_cell(mcode, r)) is not None:
                    regions[code][key(mcode, r.age, r.sex)] = c

    # --- facets: industry (学歴 第1表 industry sheets) and company size (column blocks)
    facets: Cells = {}
    wb_edu = open_book(T_EDU_AGE)
    industries: dict[str, str] = {}

    def scaled(r: MeanRow) -> list | None:
        """National distribution shape for the same sex × age, rescaled to this group's mean."""
        nat = national_means.get((r.sex, r.age))
        mq = hist_q.get((r.sex, "*", r.age))
        if not nat or not mq or r.workers < MIN_WORKERS_10:
            return None
        c = cell_histogram(nat.workers, nat.annual, nat.shoteinai, mq)
        c[0], c[1] = round(r.workers * 10), int(round(r.annual, -3))
        c[2:7] = [int(round(x * r.annual / nat.annual, -3)) for x in c[2:7]]
        return c

    for sheet in wb_edu.sheetnames:
        name = norm(sheet)
        if name.startswith("(") or name == "産業計":
            continue
        code, label = name[0], name[1:]
        industries[code] = label
        for r in parse_layout_a(wb_edu[sheet], label_col=3, first_row=13):
            if r.group == "学歴計" and (c := scaled(r)) is not None:
                facets[f"industry={code}|*|{r.age}|{r.sex}"] = c
    for size, (col, _, _) in SIZES.items():
        for r in parse_layout_a(wb_edu["産業計"], label_col=3, first_row=13, value_col=col):
            if r.group == "学歴計" and (c := scaled(r)) is not None:
                facets[f"size={size}|*|{r.age}|{r.sex}"] = c
        size_rows = parse_layout_a(open_book(T_OCC_AGE).worksheets[0], label_col=2, value_col=col)
        wb7 = open_book(T_OCC_SEX_AGE)
        for sheet, sex in (("男", "M"), ("女", "F")):
            size_rows += parse_layout_a(wb7[sheet], label_col=2, sex=sex, value_col=col)
        wb6 = open_book(T_MAJOR_SEX_AGE)
        for sheet, sex in (("男女計", "*"), ("男", "M"), ("女", "F")):
            size_rows += parse_layout_a(wb6[sheet], label_col=2, sex=sex, value_col=col)
        for r in size_rows:
            code = OCC_CODE.get(r.group or "") or MAJOR_CODE.get(r.group or "")
            if code and (c := occ_cell(code, r)) is not None:
                facets[f"size={size}|{code}|{r.age}|{r.sex}"] = c

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    write_json(OUT_DIR / "national.json", national)
    write_json(OUT_DIR / "facets.json", facets)
    for code, cells in sorted(regions.items()):
        write_json(OUT_DIR / f"region-{code}.json", cells)

    occs = LABELS["occupations"]
    meta = {
        "country": "JP",
        "currency": "JPY",
        "year": YEAR,
        "wageDefinition": {
            "ja": "年収 ＝ きまって支給する現金給与額（月額、残業代を含む）× 12 ＋ 年間賞与その他特別給与額。民営事業所（常用労働者10人以上）の一般労働者（短時間労働者を除く）。",
            "en": "Annual pay = 12 × monthly contractual cash earnings (incl. overtime) + annual bonuses. Regular (non-part-time) employees of private establishments with 10+ employees.",
        },
        "nKind": "population",
        "source": {
            "id": f"wcs{YEAR}",
            "name": {"ja": "厚生労働省「令和7年賃金構造基本統計調査」", "en": "MHLW, Basic Survey on Wage Structure 2025"},
            "url": TABLE_PAGE,
        },
        "regions": [
            {"code": PREF_CODE[ja], "label": {"ja": PREF_FULL.get(ja, ja + "県"), "en": en}}
            for ja, en in PREFS
        ],
        "occupationMajor": [{"code": f"M{k}", "label": m} for k, m in LABELS["majors"].items()],
        "occupations": [
            {"code": OCC_CODE[o["ja"]], "major": o["major"], "label": {"ja": o["ja"], "en": o["en"]}, "us": o["us"]}
            for o in occs
        ],
        "ages": AGE_BANDS,
        "educations": ["secondary", "short_tertiary", "bachelor", "graduate"],
        "facets": {
            "industry": {"label": {"ja": "業界", "en": "Industry"},
                         "values": [{"code": k, "label": {"ja": v, "en": INDUSTRY_EN.get(k, v)}} for k, v in industries.items()]},
            "size": {"label": {"ja": "企業規模", "en": "Company size"},
                     "values": [{"code": k, "label": {"ja": ja, "en": en}} for k, (_, ja, en) in SIZES.items()]},
        },
    }
    write_json(OUT_DIR / "meta.json", meta)
    write_region_summary(OUT_DIR)


if __name__ == "__main__":
    fetch()
    build()
